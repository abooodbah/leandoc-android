#!/usr/bin/env python3
"""Generates LeanDoc's test fixtures.

Run from anywhere: python3 tests/fixtures/make-fixtures.py

It writes every security fixture (SPEC section 6) into security/, every
feature fixture (SPEC section 7) into features/, a few classification
fixtures into kinds/, and the expectations both test suites read:
expected.json (reader tests) and expected.inc (core tests).

Output is deterministic: ZIP entries use a fixed 1980-01-01 timestamp, core
properties are fixed, and pseudo-random text comes from a seeded generator,
so a rerun with the same Python and zlib is byte-identical.

Needs python-docx 1.2 for the template parts and two of the documents.
"""

import io
import json
import struct
import zipfile
import zlib
from datetime import datetime, timezone
from pathlib import Path

import docx
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
FIXED_TIME = (1980, 1, 1, 0, 0, 0)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DOC_NAMESPACES = {
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "o": "urn:schemas-microsoft-com:office:office",
    "r": R_NS,
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "v": "urn:schemas-microsoft-com:vml",
    "wp14": "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "w10": "urn:schemas-microsoft-com:office:word",
    "w": W_NS,
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
}
NS_DECLS = " ".join(f'xmlns:{k}="{v}"' for k, v in DOC_NAMESPACES.items())
XML_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships/"
CT_MAIN = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "dotx": "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "docm": "application/vnd.ms-word.document.macroEnabled.main+xml",
    "dotm": "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
}
CT_WML = "application/vnd.openxmlformats-officedocument.wordprocessingml."

# ---------------------------------------------------------------------------
# Expectations, filled in by each fixture function.

EXPECTED = {}


def expect(path, kind, ok=True, error="", features=None, rtl=None, outline=None,
           text=None, purpose=""):
    """Records what the core must report for a fixture.

    outline: list of (level, text, paragraphIndex).
    text: the sanitized document's body paragraphs, in order, as plain text.
    """
    EXPECTED[path] = {
        "kind": kind,
        "ok": ok,
        "errorCode": error,
        "features": dict(features or {}),
        **({"rightToLeft": rtl} if rtl is not None else {}),
        **({"outline": [list(entry) for entry in outline]} if outline is not None else {}),
        **({"text": text} if text is not None else {}),
        "purpose": purpose,
    }


# ---------------------------------------------------------------------------
# Writing ZIP packages


def zip_bytes(entries):
    """entries: list of (name, bytes) or (name, bytes, options dict)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for entry in entries:
            name, data = entry[0], entry[1]
            options = entry[2] if len(entry) > 2 else {}
            info = zipfile.ZipInfo(name, FIXED_TIME)
            info.compress_type = options.get("method", zipfile.ZIP_DEFLATED)
            info.create_system = 0
            info.external_attr = 0
            if isinstance(data, str):
                data = data.encode("utf-8")
            archive.writestr(info, data, compresslevel=options.get("level", 9)
                             if info.compress_type == zipfile.ZIP_DEFLATED else None)
    return buffer.getvalue()


def set_flag(data, name, bit):
    """Sets a general purpose flag bit on one entry's local and central
    headers. zipfile clears flags it does not implement itself."""
    data = bytearray(data)
    encoded = name.encode("utf-8")
    for signature, flag_offset, length_offset, name_offset in (
            (b"PK\x03\x04", 6, 26, 30), (b"PK\x01\x02", 8, 28, 46)):
        position = 0
        while (position := data.find(signature, position)) >= 0:
            length = struct.unpack_from("<H", data, position + length_offset)[0]
            start = position + name_offset
            if bytes(data[start:start + length]) == encoded:
                flags = struct.unpack_from("<H", data, position + flag_offset)[0]
                struct.pack_into("<H", data, position + flag_offset, flags | bit)
            position += 4
    return bytes(data)


def write(relative, data):
    path = ROOT / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


def xml_escape(text):
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;"))


class Package:
    """An OPC package built part by part, with generated content types and
    relationships."""

    def __init__(self):
        self.parts = {}
        self.overrides = {}
        self.defaults = {
            "rels": "application/vnd.openxmlformats-package.relationships+xml",
            "xml": "application/xml",
        }
        self.rels = {}
        self.extra_entries = []

    def add(self, name, data, content_type=None):
        self.parts[name] = data
        if content_type:
            self.overrides[name] = content_type

    def rel(self, source, rel_type, target, external=False, rid=None):
        items = self.rels.setdefault(source, [])
        rid = rid or f"rId{len(items) + 1}"
        items.append((rid, rel_type, target, external))
        return rid

    def content_types(self):
        parts = [XML_DECL,
                 '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">']
        for extension, content_type in self.defaults.items():
            parts.append(f'<Default Extension="{extension}" ContentType="{content_type}"/>')
        for name, content_type in self.overrides.items():
            parts.append(f'<Override PartName="/{name}" ContentType="{content_type}"/>')
        parts.append("</Types>")
        return "".join(parts)

    @staticmethod
    def rels_name(source):
        if not source:
            return "_rels/.rels"
        folder, _, file = source.rpartition("/")
        return f"{folder}/_rels/{file}.rels" if folder else f"_rels/{file}.rels"

    def rels_xml(self, source):
        out = [XML_DECL,
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
        for rid, rel_type, target, external in self.rels[source]:
            mode = ' TargetMode="External"' if external else ""
            out.append(f'<Relationship Id="{rid}" Type="{rel_type}" '
                       f'Target="{xml_escape(target)}"{mode}/>')
        out.append("</Relationships>")
        return "".join(out)

    def build(self):
        entries = [("[Content_Types].xml", self.content_types())]
        if "" in self.rels:
            entries.append(("_rels/.rels", self.rels_xml("")))
        for name, data in self.parts.items():
            entries.append((name, data))
            if name in self.rels:
                entries.append((self.rels_name(name), self.rels_xml(name)))
        entries.extend(self.extra_entries)
        return zip_bytes(entries)


# ---------------------------------------------------------------------------
# Template parts from python-docx's default document


def python_docx_template():
    document = docx.Document()
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


TEMPLATE = python_docx_template()


def template_text(name):
    return TEMPLATE[name].decode("utf-8")


COMPACT_STYLES = XML_DECL + f'''<w:styles xmlns:w="{W_NS}">
<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:asciiTheme="minorHAnsi" w:eastAsiaTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi" w:cstheme="minorBidi"/><w:sz w:val="22"/><w:szCs w:val="22"/><w:lang w:val="en-US" w:eastAsia="en-US" w:bidi="ar-SA"/></w:rPr></w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="160" w:line="259" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>
<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>
<w:style w:type="character" w:default="1" w:styleId="DefaultParagraphFont"><w:name w:val="Default Paragraph Font"/><w:uiPriority w:val="1"/><w:semiHidden/></w:style>
<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/><w:semiHidden/><w:tblPr><w:tblInd w:w="0" w:type="dxa"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style>
<w:style w:type="numbering" w:default="1" w:styleId="NoList"><w:name w:val="No List"/><w:semiHidden/></w:style>
<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:after="240"/></w:pPr><w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi"/><w:sz w:val="56"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="360" w:after="120"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi" w:cstheme="majorBidi"/><w:color w:val="2F5496" w:themeColor="accent1" w:themeShade="BF"/><w:sz w:val="32"/><w:szCs w:val="32"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="80"/><w:outlineLvl w:val="1"/></w:pPr><w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi" w:cstheme="majorBidi"/><w:color w:val="2F5496" w:themeColor="accent1" w:themeShade="BF"/><w:sz w:val="26"/><w:szCs w:val="26"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading3"><w:name w:val="heading 3"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="160" w:after="40"/><w:outlineLvl w:val="2"/></w:pPr><w:rPr><w:color w:val="1F3763" w:themeColor="accent1" w:themeShade="7F"/><w:sz w:val="24"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading4"><w:name w:val="heading 4"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/></w:pPr><w:rPr><w:i/><w:color w:val="2F5496" w:themeColor="accent1" w:themeShade="BF"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="ListParagraph"><w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/><w:qFormat/><w:pPr><w:ind w:left="720"/><w:contextualSpacing/></w:pPr></w:style>
<w:style w:type="paragraph" w:styleId="BaseStyle"><w:name w:val="Base Style"/><w:basedOn w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:after="120"/></w:pPr><w:rPr><w:rFonts w:ascii="Georgia" w:hAnsi="Georgia"/><w:color w:val="1F4E79"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="DerivedStyle"><w:name w:val="Derived Style"/><w:basedOn w:val="BaseStyle"/><w:qFormat/><w:rPr><w:b/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="DeepStyle"><w:name w:val="Deep Style"/><w:basedOn w:val="DerivedStyle"/><w:qFormat/><w:pPr><w:jc w:val="center"/></w:pPr><w:rPr><w:i/><w:sz w:val="28"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="ChapterHeading"><w:name w:val="Chapter Heading"/><w:basedOn w:val="Heading1"/><w:qFormat/><w:rPr><w:caps/></w:rPr></w:style>
<w:style w:type="character" w:styleId="Accent"><w:name w:val="Accent"/><w:basedOn w:val="DefaultParagraphFont"/><w:qFormat/><w:rPr><w:i/><w:color w:val="C00000"/></w:rPr></w:style>
<w:style w:type="character" w:styleId="StrongAccent"><w:name w:val="Strong Accent"/><w:basedOn w:val="Accent"/><w:qFormat/><w:rPr><w:b/></w:rPr></w:style>
<w:style w:type="character" w:styleId="Hyperlink"><w:name w:val="Hyperlink"/><w:basedOn w:val="DefaultParagraphFont"/><w:rPr><w:color w:val="0563C1" w:themeColor="hyperlink"/><w:u w:val="single"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Header"><w:name w:val="header"/><w:basedOn w:val="Normal"/><w:pPr><w:tabs><w:tab w:val="center" w:pos="4680"/><w:tab w:val="right" w:pos="9360"/></w:tabs><w:spacing w:after="0"/></w:pPr></w:style>
<w:style w:type="paragraph" w:styleId="Footer"><w:name w:val="footer"/><w:basedOn w:val="Normal"/><w:pPr><w:tabs><w:tab w:val="center" w:pos="4680"/><w:tab w:val="right" w:pos="9360"/></w:tabs><w:spacing w:after="0"/></w:pPr></w:style>
<w:style w:type="paragraph" w:styleId="FootnoteText"><w:name w:val="footnote text"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="0"/></w:pPr><w:rPr><w:sz w:val="20"/></w:rPr></w:style>
<w:style w:type="character" w:styleId="FootnoteReference"><w:name w:val="footnote reference"/><w:rPr><w:vertAlign w:val="superscript"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="EndnoteText"><w:name w:val="endnote text"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="0"/></w:pPr><w:rPr><w:sz w:val="20"/></w:rPr></w:style>
<w:style w:type="character" w:styleId="EndnoteReference"><w:name w:val="endnote reference"/><w:rPr><w:vertAlign w:val="superscript"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="CommentText"><w:name w:val="annotation text"/><w:basedOn w:val="Normal"/><w:rPr><w:sz w:val="20"/></w:rPr></w:style>
<w:style w:type="character" w:styleId="CommentReference"><w:name w:val="annotation reference"/><w:rPr><w:sz w:val="16"/></w:rPr></w:style>
<w:style w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/><w:basedOn w:val="TableNormal"/><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/></w:tblBorders></w:tblPr></w:style>
</w:styles>'''

SECT_PORTRAIT = ('<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" '
                 'w:right="1440" w:bottom="1440" w:left="1440" w:header="708" '
                 'w:footer="708" w:gutter="0"/><w:cols w:space="708"/></w:sectPr>')


def core_xml(title):
    return XML_DECL + (
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f'<dc:title>{xml_escape(title)}</dc:title><dc:creator>LeanDoc fixtures</dc:creator>'
        '<cp:lastModifiedBy>LeanDoc fixtures</cp:lastModifiedBy><cp:revision>1</cp:revision>'
        '<dcterms:created xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:created>'
        '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified>'
        '</cp:coreProperties>')


APP_XML = XML_DECL + (
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
    '<Application>LeanDoc fixture generator</Application></Properties>')


def document_xml(body, sect=SECT_PORTRAIT):
    return XML_DECL + f'<w:document {NS_DECLS} mc:Ignorable="w14 wp14"><w:body>{body}{sect}</w:body></w:document>'


def settings_xml(after_proof="", after_tab_stop="", before_compat=""):
    settings = template_text("word/settings.xml")
    if after_proof:
        anchor = '<w:proofState w:spelling="clean" w:grammar="clean"/>'
        assert anchor in settings
        settings = settings.replace(anchor, anchor + after_proof)
    if after_tab_stop:
        anchor = '<w:defaultTabStop w:val="720"/>'
        assert anchor in settings
        settings = settings.replace(anchor, anchor + after_tab_stop)
    if before_compat:
        settings = settings.replace("<w:compat>", before_compat + "<w:compat>")
    return settings


def new_docx(title, body, sect=SECT_PORTRAIT, kind="docx", numbering=None,
             settings=None, styles=COMPACT_STYLES):
    package = Package()
    main = "word/document.xml"
    package.rel("", REL + "officeDocument", main)
    package.rel("", PKG_REL + "metadata/core-properties", "docProps/core.xml")
    package.rel("", REL + "extended-properties", "docProps/app.xml")
    package.add(main, document_xml(body, sect), CT_MAIN[kind])
    package.add("word/styles.xml", styles, CT_WML + "styles+xml")
    package.rel(main, REL + "styles", "styles.xml")
    package.add("word/settings.xml", settings or settings_xml(), CT_WML + "settings+xml")
    package.rel(main, REL + "settings", "settings.xml")
    package.add("word/webSettings.xml", TEMPLATE["word/webSettings.xml"], CT_WML + "webSettings+xml")
    package.rel(main, REL + "webSettings", "webSettings.xml")
    package.add("word/fontTable.xml", TEMPLATE["word/fontTable.xml"], CT_WML + "fontTable+xml")
    package.rel(main, REL + "fontTable", "fontTable.xml")
    package.add("word/theme/theme1.xml", TEMPLATE["word/theme/theme1.xml"],
                "application/vnd.openxmlformats-officedocument.theme+xml")
    package.rel(main, REL + "theme", "theme/theme1.xml")
    if numbering:
        package.add("word/numbering.xml", numbering, CT_WML + "numbering+xml")
        package.rel(main, REL + "numbering", "numbering.xml")
    package.add("docProps/core.xml", core_xml(title),
                "application/vnd.openxmlformats-package.core-properties+xml")
    package.add("docProps/app.xml", APP_XML,
                "application/vnd.openxmlformats-officedocument.extended-properties+xml")
    return package


# ---------------------------------------------------------------------------
# WordprocessingML snippets


def run(text, rpr=""):
    rpr = f"<w:rPr>{rpr}</w:rPr>" if rpr else ""
    return f'<w:r>{rpr}<w:t xml:space="preserve">{xml_escape(text)}</w:t></w:r>'


def para(content, style=None, ppr=""):
    if isinstance(content, str) and not content.startswith("<"):
        content = run(content)
    props = (f'<w:pStyle w:val="{style}"/>' if style else "") + ppr
    props = f"<w:pPr>{props}</w:pPr>" if props else ""
    return f"<w:p>{props}{content}</w:p>"


def list_para(text, num_id, level, ppr="", rpr=""):
    return para(run(text, rpr), "ListParagraph",
                f'<w:numPr><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>{ppr}')


def numbering_xml(abstracts, nums):
    return XML_DECL + f'<w:numbering xmlns:w="{W_NS}">{"".join(abstracts)}{"".join(nums)}</w:numbering>'


def level(ilvl, fmt, text, font=None, start=1, indent=None, ppr_extra=""):
    indent = indent if indent is not None else 720 * (ilvl + 1)
    rpr = (f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:hint="default"/></w:rPr>'
           if font else "")
    return (f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{xml_escape(text)}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{indent}" w:hanging="360"/>{ppr_extra}</w:pPr>{rpr}</w:lvl>')


def png(width, height, color, stripe=None):
    """A small truecolor PNG written without any imaging library."""
    rows = []
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            pixel = stripe if stripe and (x // 4 + y // 4) % 2 else color
            row.extend(pixel)
        rows.append(bytes(row))
    raw = b"".join(rows)

    def chunk(kind, data):
        return (struct.pack(">I", len(data)) + kind + data +
                struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) +
            chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))


def inline_picture(rid, docpr_id, name, cx=1828800, cy=914400, link=False):
    attribute = f'r:link="{rid}"' if link else f'r:embed="{rid}"'
    return (
        f'<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:docPr id="{docpr_id}" name="{name}"/>'
        '<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{docpr_id}" name="{name}"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip {attribute}/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
        '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def complex_field(instruction_parts, result_text):
    runs = ['<w:r><w:fldChar w:fldCharType="begin"/></w:r>']
    for part in instruction_parts:
        runs.append(f'<w:r><w:instrText xml:space="preserve">{xml_escape(part)}</w:instrText></w:r>')
    runs.append('<w:r><w:fldChar w:fldCharType="separate"/></w:r>')
    runs.append(run(result_text))
    runs.append('<w:r><w:fldChar w:fldCharType="end"/></w:r>')
    return "".join(runs)


# A small deterministic word generator for long documents.
WORDS = ("the of and to in is that for it as was with be by on not he this are or his "
         "from at which but have an they you were her all she there would their we him been "
         "has when who will more no if out so said what up its about into than them can only "
         "other new some could time these two may then do first any my now such like our over "
         "man me even most made after also did many before must through back years where much "
         "your way well down should because each just those people how too little state good "
         "very make world still own see men work long get here between both life being under "
         "never day same another know while last might us great old year off come since against "
         "go came right used take three reader document paper margin chapter section table "
         "figure river garden window station morning evening letter number pattern signal "
         "harbour mountain village library history science method result measure balance").split()


class Words:
    def __init__(self, seed):
        self.state = seed

    def next(self):
        self.state = (self.state * 1103515245 + 12345) & 0x7FFFFFFF
        return self.state

    def sentence(self, count):
        words = [WORDS[self.next() % len(WORDS)] for _ in range(count)]
        return " ".join(words).capitalize() + "."

    def paragraph(self, sentences):
        return " ".join(self.sentence(6 + self.next() % 10) for _ in range(sentences))


# ---------------------------------------------------------------------------
# Security fixtures


def security_fixtures():
    plain_body = para("A plain paragraph.")

    # zip-bomb: document.xml declares 70 MB, which inflates past the 64 MB
    # part limit. It is whitespace inside w:body, so it is even valid XML.
    package = new_docx("Zip bomb", "")
    filler = " " * (70 * 1024 * 1024)
    package.parts["word/document.xml"] = document_xml(para("Bomb.") + filler)
    write("security/zip-bomb.docx", package.build())
    expect("security/zip-bomb.docx", "docx", False, "part-too-large",
           purpose="One part declares 70 MB; inflating it would pass the 64 MB part cap.")

    package = new_docx("Many entries", plain_body)
    package.extra_entries = [(f"z/{index:04d}", b"", {"method": zipfile.ZIP_STORED})
                             for index in range(5000)]
    write("security/many-entries.docx", package.build())
    expect("security/many-entries.docx", "not-a-document", False, "too-many-entries",
           purpose="5,000 extra empty entries, past the 4,096 entry limit.")

    package = new_docx("Zip slip", plain_body)
    package.extra_entries = [("../evil.xml", "<evil/>"), ("/abs.xml", "<abs/>"),
                             ("a\\b.xml", "<b/>")]
    write("security/zip-slip.docx", package.build())
    expect("security/zip-slip.docx", "docx", False, "unsafe-entry-name",
           purpose="Entries named ../evil.xml, /abs.xml, and a\\b.xml.")

    package = new_docx("Duplicate entry", plain_body)
    package.extra_entries = [("WORD/document.xml", document_xml(para("Shadow copy.")))]
    write("security/duplicate-entry.docx", package.build())
    expect("security/duplicate-entry.docx", "docx", False, "duplicate-entry",
           purpose="word/document.xml and WORD/document.xml: equal names ignoring case.")

    full = new_docx("Truncated", plain_body).build()
    write("security/truncated.docx", full[: len(full) * 6 // 10])
    expect("security/truncated.docx", "not-a-document", False, "corrupt-zip",
           purpose="A valid package cut at 60 percent; the central directory is gone.")

    xxe = (XML_DECL.replace("?>\n", "?>") +
           '<!DOCTYPE w:document [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
           f'<w:document {NS_DECLS}><w:body><w:p><w:r><w:t>&xxe;</w:t></w:r></w:p>'
           f'{SECT_PORTRAIT}</w:body></w:document>')
    package = new_docx("XXE", "")
    package.parts["word/document.xml"] = xxe
    write("security/xxe.docx", package.build())
    expect("security/xxe.docx", "docx", False, "xml-doctype",
           purpose="document.xml declares a SYSTEM entity pointing at a local file.")

    laughs = ['<!DOCTYPE w:document [<!ENTITY lol "lol">']
    for index in range(1, 10):
        previous = "lol" if index == 1 else f"lol{index - 1}"
        laughs.append(f'<!ENTITY lol{index} "{("&" + previous + ";") * 10}">')
    laughs.append("]>")
    package = new_docx("Billion laughs", "")
    package.parts["word/document.xml"] = (
        XML_DECL + "".join(laughs) +
        f'<w:document {NS_DECLS}><w:body><w:p><w:r><w:t>&lol9;</w:t></w:r></w:p>'
        f'</w:body></w:document>')
    write("security/billion-laughs.docx", package.build())
    expect("security/billion-laughs.docx", "docx", False, "xml-doctype",
           purpose="The classic nested entity expansion.")

    nested = "<w:sdt><w:sdtContent>" * 150 + para("Deep.") + "</w:sdtContent></w:sdt>" * 150
    write("security/deep-nesting.docx", new_docx("Deep nesting", nested).build())
    expect("security/deep-nesting.docx", "docx", False, "xml-too-deep",
           purpose="300 nested elements inside w:body, past the 256 depth limit.")

    package = new_docx("Remote template", para("This document names a remote template."),
                       settings=settings_xml(after_proof='<w:attachedTemplate r:id="rId1"/>'))
    package.rel("word/settings.xml", REL + "attachedTemplate",
                "https://example.invalid/t.dotm", external=True)
    write("security/remote-template.docx", package.build())
    expect("security/remote-template.docx", "docx", features={"remote-template": 1},
           text=["This document names a remote template."],
           purpose="settings.xml attachedTemplate points at https://example.invalid/t.dotm.")

    package = new_docx("Remote image", "")
    rid = package.rel("word/document.xml", REL + "image", "https://example.invalid/a.png",
                      external=True, rid="rIdRemote")
    package.parts["word/document.xml"] = document_xml(
        para("Before the linked picture.") + para(inline_picture(rid, 1, "Linked", link=True)) +
        para("After the linked picture."))
    write("security/remote-image.docx", package.build())
    expect("security/remote-image.docx", "docx", features={"remote-image": 1},
           text=["Before the linked picture.", "", "After the linked picture."],
           purpose="A picture whose a:blip r:link points at https://example.invalid/a.png.")

    body = (para(complex_field([' INCLUDEPICTURE "https://example.invalid/p.png" \\d '],
                               "Cached picture text")) +
            para('<w:fldSimple w:instr=" INCLUDETEXT &quot;https://example.invalid/t.docx&quot; ">'
                 + run("Cached included text") + "</w:fldSimple>") +
            para("Plain text after the fields."))
    write("security/include-field.docx", new_docx("Include field", body).build())
    expect("security/include-field.docx", "docx", features={"external-content": 2},
           text=["Cached picture text", "Cached included text", "Plain text after the fields."],
           purpose="INCLUDEPICTURE (complex field, \\d) and INCLUDETEXT (simple field) with cached results.")

    body = (para(run("Before. ") + complex_field(
        [" DDE", 'AUTO c:\\\\Windows\\\\System32\\\\cmd.exe "/k calc" "!A1" '],
        "Cached DDE result") + run(" After.")))
    write("security/dde.docx", new_docx("DDE", body).build())
    expect("security/dde.docx", "docx", features={"dde-field": 1},
           text=["Before. Cached DDE result After."],
           purpose="A DDEAUTO field whose code is split across two w:instrText runs.")

    package = new_docx("Macro document", para("A macro-enabled document."), kind="docm")
    package.add("word/vbaProject.bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504 +
                b"Attribute VB_Name = \"ThisDocument\"\r\nSub AutoOpen()\r\nEnd Sub\r\n")
    package.defaults["bin"] = "application/vnd.ms-office.vbaProject"
    package.rel("word/document.xml", "http://schemas.microsoft.com/office/2006/relationships/vbaProject",
                "vbaProject.bin")
    write("security/macro.docm", package.build())
    expect("security/macro.docm", "docm", features={"macros": 1},
           text=["A macro-enabled document."],
           purpose="Macro-enabled content type with word/vbaProject.bin.")

    package = new_docx("ActiveX", "")
    image = package.rel("word/document.xml", REL + "image", "media/image1.png")
    control = package.rel("word/document.xml", REL + "control", "activeX/activeX1.xml")
    package.add("word/media/image1.png", png(16, 16, (200, 200, 200)))
    package.defaults["png"] = "image/png"
    package.add("word/activeX/activeX1.xml", XML_DECL +
                '<ax:ocx xmlns:ax="http://schemas.microsoft.com/office/2006/activeX" '
                f'xmlns:r="{R_NS}" ax:classid="{{8BD21D40-EC42-11CE-9E0D-00AA006002F3}}" '
                'ax:persistence="persistStreamInit" r:id="rId1"/>',
                "application/vnd.ms-office.activeX+xml")
    package.rel("word/activeX/activeX1.xml",
                "http://schemas.microsoft.com/office/2006/relationships/activeXControlBinary",
                "activeX1.bin")
    package.add("word/activeX/activeX1.bin", b"\x00" * 64, "application/vnd.ms-office.activeX")
    package.parts["word/document.xml"] = document_xml(
        para("A form with a check box: ") +
        para('<w:r><w:object w:dxaOrig="225" w:dyaOrig="225">'
             '<v:shape id="_x0000_i1025" type="#_x0000_t75" style="width:18pt;height:18pt">'
             f'<v:imagedata r:id="{image}" o:title=""/></v:shape>'
             f'<w:control r:id="{control}" w:name="CheckBox1" w:shapeid="_x0000_i1025"/>'
             '</w:object></w:r>'))
    write("security/activex.docx", package.build())
    expect("security/activex.docx", "docx", features={"activex": 1},
           text=["A form with a check box:", ""],
           purpose="A w:object holding a w:control backed by an activeX part.")

    package = new_docx("OLE", "")
    image = package.rel("word/document.xml", REL + "image", "media/image1.png")
    ole = package.rel("word/document.xml", REL + "oleObject", "embeddings/oleObject1.bin")
    package.add("word/media/image1.png", png(32, 24, (180, 210, 240)))
    package.defaults["png"] = "image/png"
    package.add("word/embeddings/oleObject1.bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504,
                "application/vnd.openxmlformats-officedocument.oleObject")
    package.parts["word/document.xml"] = document_xml(
        para("An embedded package follows.") +
        para('<w:r><w:object w:dxaOrig="1440" w:dyaOrig="1080">'
             '<v:shape id="_x0000_i1026" type="#_x0000_t75" style="width:72pt;height:54pt" o:ole="">'
             f'<v:imagedata r:id="{image}" o:title=""/></v:shape>'
             f'<o:OLEObject Type="Embed" ProgID="Package" ShapeID="_x0000_i1026" '
             f'DrawAspect="Icon" ObjectID="_1000000001" r:id="{ole}"/></w:object></w:r>'))
    write("security/ole.docx", package.build())
    expect("security/ole.docx", "docx", features={"ole-object": 1},
           text=["An embedded package follows.", ""],
           purpose="A w:object with o:OLEObject and an embedded payload part.")

    package = new_docx("altChunk", "")
    chunk = package.rel("word/document.xml", REL + "aFChunk", "afchunk.htm")
    package.add("word/afchunk.htm", "<html><body><script>alert('leandoc')</script>"
                "<p>Imported HTML</p></body></html>", "text/html")
    package.parts["word/document.xml"] = document_xml(
        para("Before the imported block.") + f'<w:altChunk r:id="{chunk}"/>' +
        para("After the imported block."))
    write("security/altchunk.docx", package.build())
    expect("security/altchunk.docx", "docx", features={"alt-chunk": 1},
           text=["Before the imported block.", "After the imported block."],
           purpose="w:altChunk importing an HTML part that contains a script.")

    package = new_docx("Hyperlinks", "")
    links = [("https://example.invalid/page", "HTTPS link"),
             ("mailto:reader@example.invalid", "Mail link"),
             ("javascript:alert(1)", "Script link"),
             ("file:///etc/passwd", "File link")]
    body = ""
    for target, label in links:
        rid = package.rel("word/document.xml", REL + "hyperlink", target, external=True)
        body += para(f'<w:hyperlink r:id="{rid}">' + run(label, '<w:rStyle w:val="Hyperlink"/>')
                     + "</w:hyperlink>")
    package.parts["word/document.xml"] = document_xml(body)
    write("security/hyperlinks.docx", package.build())
    expect("security/hyperlinks.docx", "docx",
           text=[label for _, label in links],
           purpose="https and mailto links survive; javascript: and file: links become plain text.")

    package = new_docx("SVG image", "")
    rid_png = package.rel("word/document.xml", REL + "image", "media/image1.png")
    rid_svg = package.rel("word/document.xml", REL + "image", "media/image2.svg")
    package.add("word/media/image1.png", png(24, 24, (30, 120, 60)))
    package.add("word/media/image2.svg", '<svg xmlns="http://www.w3.org/2000/svg">'
                "<script>alert('leandoc')</script></svg>")
    package.defaults["png"] = "image/png"
    package.defaults["svg"] = "image/svg+xml"
    picture = inline_picture(rid_png, 1, "Vector").replace(
        f'<a:blip r:embed="{rid_png}"/>',
        f'<a:blip r:embed="{rid_png}"><a:extLst><a:ext uri="{{96DAC541-7B7A-43D3-8B79-37D633B846F1}}">'
        f'<asvg:svgBlip r:embed="{rid_svg}"/></a:ext></a:extLst></a:blip>')
    package.parts["word/document.xml"] = document_xml(para("An SVG with a PNG fallback.") +
                                                      para(picture))
    write("security/svg-image.docx", package.build())
    expect("security/svg-image.docx", "docx", text=["An SVG with a PNG fallback.", ""],
           purpose="The SVG part (with a script) is dropped; the PNG fallback is kept.")

    package = new_docx("Embedded fonts", para("This document embeds a font."))
    font_table = template_text("word/fontTable.xml").replace(
        '<w:font w:name="Symbol">',
        '<w:font w:name="LeanDoc Test Sans"><w:embedRegular r:id="rId1" '
        'w:fontKey="{00000000-0000-0000-0000-000000000000}"/><w:embedBold r:id="rId2" '
        'w:fontKey="{00000000-0000-0000-0000-000000000001}"/></w:font><w:font w:name="Symbol">', 1)
    package.parts["word/fontTable.xml"] = font_table
    package.rel("word/fontTable.xml", REL + "font", "fonts/font1.odttf")
    package.rel("word/fontTable.xml", REL + "font", "fonts/font2.odttf")
    package.add("word/fonts/font1.odttf", b"\x00\x01\x00\x00" + b"\x00" * 60,
                "application/vnd.openxmlformats-officedocument.obfuscatedFont")
    package.add("word/fonts/font2.odttf", b"\x00\x01\x00\x00" + b"\x01" * 60,
                "application/vnd.openxmlformats-officedocument.obfuscatedFont")
    write("security/embedded-fonts.docx", package.build())
    expect("security/embedded-fonts.docx", "docx", features={"embedded-fonts": 2},
           text=["This document embeds a font."],
           purpose="fontTable.xml with w:embedRegular and w:embedBold font parts.")

    package = new_docx("Encrypted entry", plain_body)
    package.extra_entries = [("word/secret.xml", "<secret/>")]
    write("security/encrypted-entry.docx", set_flag(package.build(), "word/secret.xml", 0x1))
    expect("security/encrypted-entry.docx", "docx", False, "encrypted-entry",
           purpose="One ZIP entry has the traditional encryption flag set.")

    package = new_docx("Unsupported method", plain_body)
    package.extra_entries = [("word/extra.xml", "<extra/>" * 20, {"method": zipfile.ZIP_BZIP2})]
    write("security/unsupported-method.docx", package.build())
    expect("security/unsupported-method.docx", "docx", False, "corrupt-zip",
           purpose="One entry uses bzip2; only stored and deflate are read.")

    write("security/legacy.doc", compound_file([
        ("1Table", b"\x00" * 4096),
        ("WordDocument", b"\xec\xa5\xc1\x00" + b"\x00" * 4092),
    ]))
    expect("security/legacy.doc", "legacy-doc",
           purpose="An OLE compound file with WordDocument and 1Table streams.")

    info = (b"\x04\x00\x04\x00\x40\x00\x00\x00" +
            ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
             '<encryption xmlns="http://schemas.microsoft.com/office/2006/encryption">'
             '<keyData saltSize="16" blockSize="16" keyBits="256" hashSize="64" '
             'cipherAlgorithm="AES" cipherChaining="ChainingModeCBC" hashAlgorithm="SHA512" '
             'saltValue="AAAAAAAAAAAAAAAAAAAAAA=="/></encryption>').encode("ascii"))
    write("security/encrypted.docx", compound_file([
        ("EncryptionInfo", info.ljust(4096, b"\x00")),
        ("EncryptedPackage", struct.pack("<Q", 8192) + bytes((index * 7) & 0xFF for index in range(8184))),
    ]))
    expect("security/encrypted.docx", "encrypted-ooxml",
           purpose="An OLE compound file with EncryptionInfo and EncryptedPackage streams.")


def compound_file(streams):
    """A minimal version 3 compound file (512-byte sectors). Every stream is at
    least 4,096 bytes, so none lives in the mini stream."""
    sector = 512
    free, end, fat_sect = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD
    data_sectors = []
    chains = []
    next_sector = 2  # 0 = FAT, 1 = directory
    for _, data in streams:
        count = (len(data) + sector - 1) // sector
        chains.append((next_sector, count))
        data_sectors.append(data.ljust(count * sector, b"\x00"))
        next_sector += count
    fat = [fat_sect, end]
    for start, count in chains:
        fat.extend(start + offset + 1 for offset in range(count - 1))
        fat.append(end)
    assert len(fat) <= sector // 4
    fat.extend([free] * (sector // 4 - len(fat)))

    def entry(name, kind, child=free, right=free, start=0, size=0):
        encoded = name.encode("utf-16-le") + b"\x00\x00"
        return (encoded.ljust(64, b"\x00") + struct.pack("<HBB", len(encoded), kind, 1) +
                struct.pack("<III", free, right, child) + b"\x00" * 16 + b"\x00" * 4 +
                b"\x00" * 16 + struct.pack("<IQ", start, size))

    entries = [entry("Root Entry", 5, child=1, start=end)]
    for index, ((name, data), (start, _)) in enumerate(zip(streams, chains)):
        right = index + 2 if index + 1 < len(streams) else free
        entries.append(entry(name, 2, right=right, start=start, size=len(data)))
    directory = b"".join(entries).ljust(sector, b"\x00")

    header = (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16 +
              struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6) + b"\x00" * 6 +
              struct.pack("<IIIIIIIII", 0, 1, 1, 0, 4096, end, 0, end, 0))
    header += struct.pack("<I", 0) + struct.pack("<I", free) * 108
    assert len(header) == sector
    return header + struct.pack(f"<{sector // 4}I", *fat) + directory + b"".join(data_sectors)


# ---------------------------------------------------------------------------
# Classification fixtures


def kind_fixtures():
    write("kinds/plain.rtf", b"{\\rtf1\\ansi\\deff0 {\\fonttbl {\\f0 Times New Roman;}}"
                             b"\\f0\\fs24 An RTF document.\\par}")
    expect("kinds/plain.rtf", "rtf", purpose="A short RTF file.")

    write("kinds/minimal.odt", zip_bytes([
        ("mimetype", "application/vnd.oasis.opendocument.text", {"method": zipfile.ZIP_STORED}),
        ("content.xml", XML_DECL + '<office:document-content xmlns:office="urn:oasis:names:tc:'
                                   'opendocument:xmlns:office:1.0" office:version="1.3"/>'),
    ]))
    expect("kinds/minimal.odt", "odt", purpose="A ZIP whose mimetype entry names ODF text.")

    write("kinds/spreadsheet.xlsx", zip_bytes([
        ("[Content_Types].xml", XML_DECL +
         '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
         '<Default Extension="xml" ContentType="application/xml"/>'
         '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
         'officedocument.spreadsheetml.sheet.main+xml"/></Types>'),
        ("xl/workbook.xml", XML_DECL + '<workbook xmlns="http://schemas.openxmlformats.org/'
                                       'spreadsheetml/2006/main"/>'),
    ]))
    expect("kinds/spreadsheet.xlsx", "not-a-document", False, "missing-main-part",
           purpose="OOXML, but a spreadsheet: no Word main part.")

    write("kinds/text-named.docx", b"This is plain text with a .docx name.\n")
    expect("kinds/text-named.docx", "not-a-document", False, "not-a-zip",
           purpose="Plain text with a Word extension.")

    package = new_docx("Template", para("A Word template.", "Heading1") + para("Body."),
                       kind="dotx")
    write("kinds/template.dotx", package.build())
    expect("kinds/template.dotx", "dotx", outline=[(1, "A Word template.", 0)],
           text=["A Word template.", "Body."], purpose="A .dotx template main part.")

    package = new_docx("Macro template", para("A macro-enabled template."), kind="dotm")
    write("kinds/template.dotm", package.build())
    expect("kinds/template.dotm", "dotm", features={"macros": 1},
           text=["A macro-enabled template."],
           purpose="A .dotm content type without a macro part still counts as macros.")

    strict_w = "http://purl.oclc.org/ooxml/wordprocessingml/main"
    strict_r = "http://purl.oclc.org/ooxml/officeDocument/relationships"
    package = Package()
    package.rel("", strict_r + "/officeDocument", "word/document.xml")
    package.add("word/document.xml", XML_DECL +
                f'<w:document xmlns:w="{strict_w}" xmlns:r="{strict_r}" w:conformance="strict">'
                '<w:body><w:p><w:pPr><w:outlineLvl w:val="0"/></w:pPr><w:r><w:t>Strict heading'
                '</w:t></w:r></w:p><w:p><w:r><w:t>Strict body.</w:t></w:r></w:p></w:body>'
                '</w:document>', CT_MAIN["docx"])
    write("kinds/strict.docx", package.build())
    expect("kinds/strict.docx", "docx", outline=[(1, "Strict heading", 0)],
           text=["Strict heading", "Strict body."],
           purpose="Strict OOXML namespaces, matched by URI.")

    write("kinds/empty.zip", zip_bytes([]))
    expect("kinds/empty.zip", "not-a-document", False, "missing-main-part",
           purpose="An empty ZIP archive.")


# ---------------------------------------------------------------------------
# Feature fixtures


def feature_fixtures():
    styles_fixture()
    theme_fixture()
    numbering_fixture()
    bullets_fixture()
    tables_fixture()
    images_fixture()
    headers_footers_fixture()
    notes_fixture()
    notes_superscript_fixture()
    hyperlinks_fixture()
    breaks_fixture()
    arabic_fixtures()
    not_displayed_fixture()
    malformed_equation_fixture()
    drawn_content_fixture()
    tracked_changes_fixture()
    comments_fixture()
    long_fixture()


def styles_fixture():
    body = (para("Styles and inheritance", "Title") +
            para("Chapter one", "ChapterHeading") +
            para("Base style paragraph in Georgia and blue.", "BaseStyle") +
            para("Derived style adds bold.", "DerivedStyle") +
            para("Deep style adds italics, a larger size, and centring.", "DeepStyle") +
            para("Heading by name only", "Heading4") +
            para(run("Normal text with ") + run("an accent", '<w:rStyle w:val="Accent"/>') +
                 run(" and ") + run("a strong accent", '<w:rStyle w:val="StrongAccent"/>') +
                 run(" character style.")) +
            para("Direct formatting", ppr='<w:outlineLvl w:val="1"/>') +
            para("Body text level nine", "Heading1", '<w:outlineLvl w:val="9"/>'))
    write("features/styles.docx", new_docx("Styles", body).build())
    expect("features/styles.docx", "docx", rtl=False,
           outline=[(1, "Chapter one", 1), (4, "Heading by name only", 5),
                    (2, "Direct formatting", 7)],
           text=["Styles and inheritance", "Chapter one",
                 "Base style paragraph in Georgia and blue.", "Derived style adds bold.",
                 "Deep style adds italics, a larger size, and centring.",
                 "Heading by name only",
                 "Normal text with an accent and a strong accent character style.",
                 "Direct formatting", "Body text level nine"],
           purpose="A basedOn chain of three paragraph styles, a heading style inherited "
                   "through basedOn, a heading level from the built-in name, character "
                   "styles based on each other, direct outlineLvl, and outlineLvl 9 (body).")


def theme_fixture():
    body = (para("Theme fonts and colours", "Heading1") +
            para(run("Major theme font. ", '<w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi"/>') +
                 run("Minor theme font.", '<w:rFonts w:asciiTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi"/>')) +
            para(run("Accent 1. ", '<w:color w:val="4F81BD" w:themeColor="accent1"/>') +
                 run("Accent 2. ", '<w:color w:val="C0504D" w:themeColor="accent2"/>') +
                 run("Accent 6, darker.", '<w:color w:val="F79646" w:themeColor="accent6" w:themeShade="BF"/>')) +
            para(run("Text 1 on a theme shaded paragraph."),
                 ppr='<w:shd w:val="clear" w:color="auto" w:fill="DBE5F1" w:themeFill="accent1" w:themeFillTint="33"/>'))
    write("features/theme.docx", new_docx("Theme", body).build())
    expect("features/theme.docx", "docx", rtl=False,
           outline=[(1, "Theme fonts and colours", 0)],
           purpose="Runs in the major and minor theme fonts, theme colours with shade, "
                   "and theme paragraph shading.")


def numbering_fixture():
    abstract = ('<w:abstractNum w:abstractNumId="0"><w:multiLevelType w:val="multilevel"/>' +
                level(0, "decimal", "%1.") + level(1, "lowerLetter", "%2)") +
                level(2, "lowerRoman", "%3.") + "</w:abstractNum>")
    nums = ['<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>',
            '<w:num w:numId="2"><w:abstractNumId w:val="0"/><w:lvlOverride w:ilvl="0">'
            '<w:startOverride w:val="1"/></w:lvlOverride></w:num>']
    body = (para("Multi-level numbering", "Heading1") +
            list_para("First item", 1, 0) + list_para("Nested letter a", 1, 1) +
            list_para("Deeper roman i", 1, 2) + list_para("Deeper roman ii", 1, 2) +
            list_para("Nested letter b", 1, 1) + list_para("Second item", 1, 0) +
            para("A paragraph between the lists.") +
            list_para("Restarted at one", 2, 0) + list_para("Restarted nested a", 2, 1) +
            list_para("Restarted two", 2, 0))
    write("features/numbering.docx",
          new_docx("Numbering", body, numbering=numbering_xml([abstract], nums)).build())
    expect("features/numbering.docx", "docx", rtl=False,
           outline=[(1, "Multi-level numbering", 0)],
           purpose="Three levels (decimal, lowerLetter, lowerRoman) and a second list "
                   "restarted with startOverride.")


def bullets_fixture():
    bullets = [("Symbol", "\uf0b7"), ("Wingdings", "\uf0a7"), ("Wingdings", "\uf0d8"),
               ("Wingdings", "\uf0fc"), ("Courier New", "o")]
    abstracts = []
    nums = []
    body = para("Bullets in symbol fonts", "Heading1")
    for index, (font, character) in enumerate(bullets):
        abstract_id = 10 + index
        abstracts.append(f'<w:abstractNum w:abstractNumId="{abstract_id}">'
                         '<w:multiLevelType w:val="hybridMultilevel"/>' +
                         level(0, "bullet", character, font=font) + "</w:abstractNum>")
        nums.append(f'<w:num w:numId="{abstract_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>')
        code = f"U+{ord(character):04X}"
        body += list_para(f"{font} bullet {code}", abstract_id, 0)
        body += list_para(f"Second {font} bullet", abstract_id, 0)
    write("features/bullets.docx",
          new_docx("Bullets", body, numbering=numbering_xml(abstracts, nums)).build())
    expect("features/bullets.docx", "docx", rtl=False,
           outline=[(1, "Bullets in symbol fonts", 0)],
           purpose="Bullets in Symbol (U+F0B7), Wingdings (U+F0A7, U+F0D8, U+F0FC), and "
                   "Courier New (o).")


def cell(text, width, props=""):
    return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{props}</w:tcPr>'
            f'{para(text) if text is not None else "<w:p/>"}</w:tc>')


def tables_fixture():
    shade = '<w:shd w:val="clear" w:color="auto" w:fill="D9E2F3"/>'
    table = (
        '<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="9000" w:type="dxa"/>'
        '<w:tblBorders><w:top w:val="double" w:sz="4" w:space="0" w:color="1F4E79"/>'
        '<w:left w:val="single" w:sz="8" w:space="0" w:color="1F4E79"/>'
        '<w:bottom w:val="double" w:sz="4" w:space="0" w:color="1F4E79"/>'
        '<w:right w:val="single" w:sz="8" w:space="0" w:color="1F4E79"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/>'
        '<w:insideV w:val="dashed" w:sz="4" w:space="0" w:color="8EAADB"/></w:tblBorders>'
        '</w:tblPr><w:tblGrid><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/><w:gridCol w:w="3000"/></w:tblGrid>'
        '<w:tr>' + cell("Spans two columns", 6000, '<w:gridSpan w:val="2"/>' + shade) +
        cell("Header three", 3000, shade) + '</w:tr>'
        '<w:tr>' + cell("Spans two rows", 3000, '<w:vMerge w:val="restart"/><w:vAlign w:val="center"/>') +
        cell("Row two, B", 3000) +
        cell("Shaded yellow", 3000, '<w:shd w:val="clear" w:color="auto" w:fill="FFF2CC"/>') + '</w:tr>'
        '<w:tr>' + cell(None, 3000, '<w:vMerge/>') + cell("Row three, B", 3000) +
        cell("Thick red border", 3000, '<w:tcBorders><w:top w:val="single" w:sz="18" w:space="0" '
             'w:color="C00000"/><w:bottom w:val="single" w:sz="18" w:space="0" w:color="C00000"/>'
             '</w:tcBorders>') + '</w:tr></w:tbl>')
    body = para("Merged cells, borders, and shading", "Heading1") + table + para("After the table.")
    write("features/tables.docx", new_docx("Tables", body).build())
    expect("features/tables.docx", "docx", rtl=False,
           outline=[(1, "Merged cells, borders, and shading", 0)],
           text=["Merged cells, borders, and shading", "Spans two columns", "Header three",
                 "Spans two rows", "Row two, B", "Shaded yellow", "", "Row three, B",
                 "Thick red border", "After the table."],
           purpose="gridSpan and vMerge, table and cell borders, and cell shading.")


def fixed_python_docx(document, title):
    properties = document.core_properties
    properties.title = title
    properties.author = "LeanDoc fixtures"
    properties.last_modified_by = "LeanDoc fixtures"
    properties.revision = 1
    properties.created = datetime(2026, 1, 1, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        entries = [(name, archive.read(name)) for name in archive.namelist()]
    return zip_bytes(entries)


def images_fixture():
    document = docx.Document()
    document.add_heading("Inline images", 1)
    document.add_paragraph("A blue picture follows.")
    document.add_picture(io.BytesIO(png(64, 32, (40, 90, 200), (250, 250, 250))), width=Inches(2))
    document.add_paragraph("A green picture, inline with text: ")
    paragraph = document.paragraphs[-1]
    paragraph.add_run().add_picture(io.BytesIO(png(16, 16, (20, 160, 60))), width=Pt(24))
    paragraph.add_run(" and text after it.")
    document.add_picture(io.BytesIO(png(48, 48, (200, 60, 40), (255, 220, 0))), width=Inches(1))
    write("features/images.docx", fixed_python_docx(document, "Images"))
    expect("features/images.docx", "docx", rtl=False, outline=[(1, "Inline images", 0)],
           text=["Inline images", "A blue picture follows.", "",
                 "A green picture, inline with text: and text after it.", ""],
           purpose="Three PNG pictures added through python-docx, one inline in a text run.")


def headers_footers_fixture():
    package = new_docx("Headers and footers", "",
                       settings=settings_xml(after_tab_stop="<w:evenAndOddHeaders/>"))
    refs = []
    for kind, label in (("first", "First page"), ("default", "Odd page"), ("even", "Even page")):
        for part in ("header", "footer"):
            name = f"word/{part}-{kind}.xml"
            tag = "hdr" if part == "header" else "ftr"
            style = "Header" if part == "header" else "Footer"
            content = para(run(f"{label} {part}") + (
                run(" | page ") + '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
                if part == "footer" else ""), style)
            package.add(name, XML_DECL + f'<w:{tag} {NS_DECLS}>{content}</w:{tag}>',
                        CT_WML + f"{part}+xml")
            rid = package.rel("word/document.xml", REL + part, f"{part}-{kind}.xml")
            refs.append(f'<w:{part}Reference w:type="{kind}" r:id="{rid}"/>')
    sect = SECT_PORTRAIT.replace("<w:sectPr>", "<w:sectPr>" + "".join(refs)).replace(
        "</w:sectPr>", "<w:titlePg/></w:sectPr>")
    body = ""
    for page in range(1, 5):
        body += para(f"Page {page}", "Heading1") + para(f"Body text on page {page}.")
        if page < 4:
            body += para('<w:r><w:br w:type="page"/></w:r>')
    package.parts["word/document.xml"] = document_xml(body, sect)
    write("features/headers-footers.docx", package.build())
    expect("features/headers-footers.docx", "docx", rtl=False,
           outline=[(1, f"Page {page}", (page - 1) * 3) for page in range(1, 5)],
           purpose="Different first page and even/odd headers and footers over four pages, "
                   "with a PAGE field in each footer.")


def notes_fixture():
    package = new_docx("Footnotes and endnotes", "", settings=settings_xml(
        before_compat='<w:footnotePr><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>'
                      '<w:endnotePr><w:endnote w:id="-1"/><w:endnote w:id="0"/></w:endnotePr>'))

    def notes(kind, items):
        tag = kind  # footnote or endnote
        style = "FootnoteText" if kind == "footnote" else "EndnoteText"
        reference = "FootnoteReference" if kind == "footnote" else "EndnoteReference"
        out = [XML_DECL, f'<w:{tag}s {NS_DECLS}>',
               f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0"/></w:pPr>'
               '<w:r><w:separator/></w:r></w:p></w:' + tag + '>',
               f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0"/>'
               '</w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:' + tag + '>']
        for note_id, text in items:
            out.append(f'<w:{tag} w:id="{note_id}"><w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
                       f'<w:r><w:rPr><w:rStyle w:val="{reference}"/></w:rPr><w:{tag}Ref/></w:r>'
                       f'{run(" " + text)}</w:p></w:{tag}>')
        out.append(f"</w:{tag}s>")
        return "".join(out)

    package.add("word/footnotes.xml", notes("footnote", [(1, "First footnote text."),
                                                         (2, "Second footnote text.")]),
                CT_WML + "footnotes+xml")
    package.rel("word/document.xml", REL + "footnotes", "footnotes.xml")
    package.add("word/endnotes.xml", notes("endnote", [(1, "The only endnote.")]),
                CT_WML + "endnotes+xml")
    package.rel("word/document.xml", REL + "endnotes", "endnotes.xml")

    def ref(kind, note_id):
        style = "FootnoteReference" if kind == "footnote" else "EndnoteReference"
        return (f'<w:r><w:rPr><w:rStyle w:val="{style}"/></w:rPr>'
                f'<w:{kind}Reference w:id="{note_id}"/></w:r>')

    body = (para("Notes", "Heading1") +
            para(run("A sentence with a footnote.") + ref("footnote", 1) +
                 run(" Another with a second footnote.") + ref("footnote", 2)) +
            para(run("A sentence with an endnote.") + ref("endnote", 1)))
    package.parts["word/document.xml"] = document_xml(body)
    write("features/notes.docx", package.build())
    expect("features/notes.docx", "docx", rtl=False, outline=[(1, "Notes", 0)],
           text=["Notes", "A sentence with a footnote. Another with a second footnote.",
                 "A sentence with an endnote."],
           purpose="Two footnotes and one endnote, hand-written footnotes.xml and endnotes.xml.")


def notes_superscript_fixture():
    # Word writes some note references with direct superscript formatting
    # instead of the FootnoteReference style. docx-preview 0.4.1 renders such
    # a run's children twice, which numbered the first footnote "2".
    package = new_docx("Footnotes and endnotes with direct superscript", "", settings=settings_xml(
        before_compat='<w:footnotePr><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>'
                      '<w:endnotePr><w:endnote w:id="-1"/><w:endnote w:id="0"/></w:endnotePr>'))

    def notes(tag, items):
        out = [XML_DECL, f'<w:{tag}s {NS_DECLS}>',
               f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>',
               f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/>'
               f'</w:r></w:p></w:{tag}>']
        for note_id, text in items:
            out.append(f'<w:{tag} w:id="{note_id}"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/>'
                       f'</w:rPr><w:{tag}Ref/></w:r>{run(" " + text)}</w:p></w:{tag}>')
        out.append(f"</w:{tag}s>")
        return "".join(out)

    package.add("word/footnotes.xml", notes("footnote", [(1, "Footnote one, direct superscript."),
                                                         (2, "Footnote two, styled."),
                                                         (3, "Footnote three, direct superscript.")]),
                CT_WML + "footnotes+xml")
    package.rel("word/document.xml", REL + "footnotes", "footnotes.xml")
    package.add("word/endnotes.xml", notes("endnote", [(1, "Endnote one, direct superscript."),
                                                       (2, "Endnote two, styled.")]),
                CT_WML + "endnotes+xml")
    package.rel("word/document.xml", REL + "endnotes", "endnotes.xml")

    def direct(kind, note_id):
        return (f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
                f'<w:{kind}Reference w:id="{note_id}"/></w:r>')

    def styled(kind, note_id):
        style = "FootnoteReference" if kind == "footnote" else "EndnoteReference"
        return (f'<w:r><w:rPr><w:rStyle w:val="{style}"/></w:rPr>'
                f'<w:{kind}Reference w:id="{note_id}"/></w:r>')

    superscript_text = '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>th</w:t></w:r>'
    body = (para("Mixed note references", "Heading1") +
            para(run("First footnote.") + direct("footnote", 1) +
                 run(" Second footnote.") + styled("footnote", 2) +
                 run(" The 4") + superscript_text + run(" day.") +
                 run(" Third footnote.") + direct("footnote", 3)) +
            para(run("First endnote.") + direct("endnote", 1) +
                 run(" Second endnote.") + styled("endnote", 2)))
    package.parts["word/document.xml"] = document_xml(body)
    write("features/notes-superscript.docx", package.build())
    expect("features/notes-superscript.docx", "docx", rtl=False,
           outline=[(1, "Mixed note references", 0)],
           text=["Mixed note references",
                 "First footnote. Second footnote. The 4th day. Third footnote.",
                 "First endnote. Second endnote."],
           purpose="Footnote and endnote references with direct superscript formatting mixed "
                   "with styled ones; each must be numbered once, in order.")


def hyperlinks_fixture():
    package = new_docx("Links", "")
    rid = package.rel("word/document.xml", REL + "hyperlink", "https://example.invalid/docs",
                      external=True)
    mail = package.rel("word/document.xml", REL + "hyperlink", "mailto:team@example.invalid",
                       external=True)
    style = '<w:rStyle w:val="Hyperlink"/>'
    body = (para("Links", "Heading1") +
            para(run("Visit ") + f'<w:hyperlink r:id="{rid}" w:history="1">' + run("the docs", style) +
                 "</w:hyperlink>" + run(" or write to ") + f'<w:hyperlink r:id="{mail}">' +
                 run("the team", style) + "</w:hyperlink>" + run(".")) +
            para(run("Jump to ") + '<w:hyperlink w:anchor="target">' + run("the target", style) +
                 "</w:hyperlink>" + run(" below.")) +
            para('<w:r><w:br w:type="page"/></w:r>') +
            para('<w:bookmarkStart w:id="0" w:name="target"/>' + run("The target paragraph.") +
                 '<w:bookmarkEnd w:id="0"/>', "Heading2"))
    package.parts["word/document.xml"] = document_xml(body)
    write("features/hyperlinks.docx", package.build())
    expect("features/hyperlinks.docx", "docx", rtl=False,
           outline=[(1, "Links", 0), (2, "The target paragraph.", 4)],
           text=["Links", "Visit the docs or write to the team.", "Jump to the target below.",
                 "", "The target paragraph."],
           purpose="External https and mailto links and an internal bookmark link.")


def breaks_fixture():
    portrait_end = ('<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" '
                    'w:right="1440" w:bottom="1440" w:left="1440" w:header="708" w:footer="708" '
                    'w:gutter="0"/><w:cols w:space="708"/></w:sectPr>')
    landscape = ('<w:sectPr><w:type w:val="nextPage"/><w:pgSz w:w="16838" w:h="11906" '
                 'w:orient="landscape"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" '
                 'w:left="1440" w:header="708" w:footer="708" w:gutter="0"/>'
                 '<w:cols w:space="708"/></w:sectPr>')
    body = (para("Portrait section", "Heading1") + para("First page of the portrait section.") +
            para(run("Text before a page break.") + '<w:r><w:br w:type="page"/></w:r>' +
                 run("Text after the page break.")) +
            para("Last portrait paragraph.", ppr=portrait_end) +
            para("Landscape section", "Heading1") + para("This section is landscape."))
    write("features/breaks.docx", new_docx("Breaks", body, sect=landscape).build())
    expect("features/breaks.docx", "docx", rtl=False,
           outline=[(1, "Portrait section", 0), (1, "Landscape section", 4)],
           text=["Portrait section", "First page of the portrait section.",
                 "Text before a page break. Text after the page break.",
                 "Last portrait paragraph.", "Landscape section", "This section is landscape."],
           purpose="A page break inside a paragraph and a section break from portrait to "
                   "landscape.")


ARABIC_RPR = '<w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:hint="cs"/><w:rtl/><w:lang w:bidi="ar-SA"/>'


def arabic_para(text, style=None, extra_ppr=""):
    return para(run(text, ARABIC_RPR), style, "<w:bidi/>" + extra_ppr)


def arabic_fixtures():
    title = "مستند عربي"
    lines = ["هذه فقرة عربية تكتب من اليمين إلى اليسار.",
             "الفقرة الثانية تحتوي على نص أطول قليلا لاختبار الالتفاف والمحاذاة.",
             "النهاية."]
    body = arabic_para(title, "Heading1") + "".join(arabic_para(line) for line in lines)
    write("features/arabic.docx", new_docx("Arabic", body).build())
    expect("features/arabic.docx", "docx", rtl=True, outline=[(1, title, 0)],
           text=[title] + lines,
           purpose="Arabic paragraphs with w:bidi, runs with w:rtl, and the cs font hint.")

    abstract = ('<w:abstractNum w:abstractNumId="20"><w:multiLevelType w:val="hybridMultilevel"/>' +
                level(0, "decimal", "%1.", ppr_extra="<w:bidi/>") +
                level(1, "arabicAlpha", "%2.", ppr_extra="<w:bidi/>") + "</w:abstractNum>")
    nums = ['<w:num w:numId="20"><w:abstractNumId w:val="20"/></w:num>']
    items = ["البند الأول", "البند الثاني", "بند فرعي", "البند الثالث"]
    body = arabic_para("قائمة مرقمة", "Heading2")
    for index, item in enumerate(items):
        body += list_para(item, 20, 1 if index == 2 else 0, "<w:bidi/>", ARABIC_RPR)
    write("features/arabic-list.docx",
          new_docx("Arabic list", body, numbering=numbering_xml([abstract], nums)).build())
    expect("features/arabic-list.docx", "docx", rtl=True, outline=[(2, "قائمة مرقمة", 0)],
           text=["قائمة مرقمة"] + items,
           purpose="An Arabic numbered list: decimal and arabicAlpha levels on bidi paragraphs.")

    def arabic_cell(text):
        return (f'<w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr>'
                f'{arabic_para(text)}</w:tc>')

    table = ('<w:tbl><w:tblPr><w:bidiVisual/><w:tblStyle w:val="TableGrid"/>'
             '<w:tblW w:w="9000" w:type="dxa"/></w:tblPr><w:tblGrid><w:gridCol w:w="3000"/>'
             '<w:gridCol w:w="3000"/><w:gridCol w:w="3000"/></w:tblGrid>'
             '<w:tr>' + arabic_cell("الاسم") + arabic_cell("المدينة") + arabic_cell("العمر") + '</w:tr>'
             '<w:tr>' + arabic_cell("سارة") + arabic_cell("دمشق") + arabic_cell("٣٠") + '</w:tr>'
             '</w:tbl>')
    body = arabic_para("جدول", "Heading2") + table + arabic_para("بعد الجدول.")
    write("features/arabic-table.docx", new_docx("Arabic table", body).build())
    expect("features/arabic-table.docx", "docx", rtl=True, outline=[(2, "جدول", 0)],
           text=["جدول", "الاسم", "المدينة", "العمر", "سارة", "دمشق", "٣٠", "بعد الجدول."],
           purpose="A table with w:bidiVisual: the first column is on the right.")

    mixed = [para("Mixed directions", "Heading1"),
             para("An English paragraph mentions العربية in the middle."),
             para(run("فقرة عربية ", ARABIC_RPR) + run("with English words") +
                  run(" في الوسط.", ARABIC_RPR), ppr="<w:bidi/>"),
             para("Numbers 2026 and ١٤٤٧ side by side.")]
    write("features/mixed-direction.docx", new_docx("Mixed", "".join(mixed)).build())
    expect("features/mixed-direction.docx", "docx", rtl=False,
           outline=[(1, "Mixed directions", 0)],
           text=["Mixed directions", "An English paragraph mentions العربية in the middle.",
                 "فقرة عربية with English words في الوسط.",
                 "Numbers 2026 and ١٤٤٧ side by side."],
           purpose="English and Arabic in one document and in one paragraph; the first "
                   "paragraph is left to right.")


def not_displayed_fixture():
    package = new_docx("Not displayed", "")
    chart = package.rel("word/document.xml", REL + "chart", "charts/chart1.xml")
    package.add("word/charts/chart1.xml", XML_DECL +
                '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart>'
                '<c:plotArea><c:layout/></c:plotArea></c:chart></c:chartSpace>',
                "application/vnd.openxmlformats-officedocument.drawingml.chart+xml")
    dgm = {}
    for part, content_type in (("data", "diagramData"), ("layout", "diagramLayout"),
                               ("quickStyle", "diagramStyle"), ("colors", "diagramColors")):
        name = f"diagrams/{part}1.xml"
        package.add(f"word/{name}", XML_DECL +
                    f'<dgm:{"dataModel" if part == "data" else "styleDef"} '
                    'xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"/>',
                    f"application/vnd.openxmlformats-officedocument.drawingml.{content_type}+xml")
        dgm[part] = package.rel("word/document.xml", REL + content_type, name)

    def anchor_frame(docpr_id, name, graphic, cx=2743200, cy=1600200):
        return (f'<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
                f'<wp:extent cx="{cx}" cy="{cy}"/><wp:docPr id="{docpr_id}" name="{name}"/>'
                f'<wp:cNvGraphicFramePr/><a:graphic>{graphic}</a:graphic></wp:inline>'
                '</w:drawing></w:r>')

    # docx-preview 0.4.1 draws VML text boxes and most equations, so this
    # fixture uses the forms it does not draw: a DrawingML text box with no
    # VML fallback, an equation with m:sSubSup, and a shape whose fallback is
    # a v:roundrect. drawn-content.docx holds the forms it does draw.
    text_box = '<w:r>' + wps_text_box(1, "Text inside a text box.") + '</w:r>'

    shape = (
        '<w:r><mc:AlternateContent><mc:Choice Requires="wps">' + wps_rectangle(2) +
        '</mc:Choice><mc:Fallback><w:pict><v:roundrect id="Rectangle 2" '
        'style="width:72pt;height:36pt" fillcolor="#4472c4"/></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>')

    equation = ('<m:oMathPara><m:oMath><m:sSubSup><m:e><m:r><m:t>x</m:t></m:r></m:e>'
                '<m:sub><m:r><m:t>i</m:t></m:r></m:sub><m:sup><m:r><m:t>2</m:t></m:r></m:sup>'
                '</m:sSubSup></m:oMath></m:oMathPara>')

    chart_frame = anchor_frame(3, "Chart 3", '<a:graphicData uri="http://schemas.openxmlformats.org/'
                               f'drawingml/2006/chart"><c:chart r:id="{chart}"/></a:graphicData>')
    smartart = anchor_frame(4, "Diagram 4", '<a:graphicData uri="http://schemas.openxmlformats.org/'
                            'drawingml/2006/diagram"><dgm:relIds '
                            f'r:dm="{dgm["data"]}" r:lo="{dgm["layout"]}" r:qs="{dgm["quickStyle"]}" '
                            f'r:cs="{dgm["colors"]}"/></a:graphicData>')

    body = (para("Content LeanDoc does not display", "Heading1") +
            para(run("A text box: ") + text_box) +
            para(equation) +
            para(run("A chart: ") + chart_frame) +
            para(run("A SmartArt graphic: ") + smartart) +
            para(run("A shape: ") + shape) +
            para("The end."))
    package.parts["word/document.xml"] = document_xml(body)
    write("features/not-displayed.docx", package.build())
    expect("features/not-displayed.docx", "docx", rtl=False,
           features={"text-box": 1, "equation": 1, "chart": 1, "smartart": 1, "shape": 1},
           outline=[(1, "Content LeanDoc does not display", 0)],
           text=["Content LeanDoc does not display", "A text box:", "[Equation]", "A chart:",
                 "A SmartArt graphic:", "A shape:", "The end."],
           purpose="One of each not-displayed feature in a form docx-preview 0.4.1 does not "
                   "draw: text box, equation, chart, SmartArt, and shape, for the fidelity bar.")


def wps_text_box(docpr_id, text):
    """A floating DrawingML text box, as Word 2010 and later write it."""
    return (
        '<w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" '
        f'relativeHeight="{251659264 + docpr_id}" behindDoc="0" locked="0" layoutInCell="1" '
        'allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column">'
        '<wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph">'
        '<wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="2286000" cy="457200"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="bothSides"/>'
        f'<wp:docPr id="{docpr_id}" name="Text Box {docpr_id}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/>'
        '<a:ext cx="2286000" cy="457200"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '</wps:spPr><wps:txbx><w:txbxContent>' + para(text) +
        '</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>'
        '</wp:anchor></w:drawing>')


def vml_text_box(shape_id, text, style="width:180pt;height:36pt"):
    return (f'<w:pict><v:shape id="Text Box {shape_id}" type="#_x0000_t202" style="{style}">'
            '<v:textbox><w:txbxContent>' + para(text) +
            '</w:txbxContent></v:textbox></v:shape></w:pict>')


def wps_rectangle(docpr_id):
    return (
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="914400" cy="457200"/><wp:docPr id="{docpr_id}" name="Rectangle {docpr_id}"/>'
        '<wp:cNvGraphicFramePr/><a:graphic>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="457200"/>'
        '</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="4472C4"/>'
        '</a:solidFill></wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline>'
        '</w:drawing>')


def malformed_equation_fixture():
    # docx-preview 0.4.1 throws on a delimiter without m:dPr, which would blank
    # the whole document. The core replaces such an equation with a marker.
    package = new_docx("Malformed equation", "")
    equation = ('<m:oMath><m:d><m:e><m:r><m:t>a+b</m:t></m:r></m:e></m:d></m:oMath>')
    body = (para("Malformed equation", "Heading1") +
            para("Text before the equation.") +
            para(equation) +
            para("Text after the equation."))
    package.parts["word/document.xml"] = document_xml(body)
    write("features/malformed-equation.docx", package.build())
    expect("features/malformed-equation.docx", "docx", rtl=False,
           features={"equation": 1},
           outline=[(1, "Malformed equation", 0)],
           text=["Malformed equation", "Text before the equation.", "[Equation]",
                 "Text after the equation."],
           purpose="An equation docx-preview 0.4.1 throws on (m:d without m:dPr). The "
                   "document must still render, with a marker where the equation was.")


def drawn_content_fixture():
    """Text boxes, a shape, and equations in the forms docx-preview 0.4.1 draws.

    The core reports none of them as not displayed. Word writes a DrawingML
    text box with a VML fallback; docx-preview supports no mc:Choice and
    draws the fallback. It renders OMML fractions, scripts, radicals, n-ary
    operators, and delimiters as MathML.
    """
    def m_run(text):
        return f"<m:r><m:t>{xml_escape(text)}</m:t></m:r>"

    with_fallback = (
        '<w:r><mc:AlternateContent><mc:Choice Requires="wps">' +
        wps_text_box(1, "Text in a text box with a fallback.") +
        '</mc:Choice><mc:Fallback>' +
        vml_text_box(1, "Text in a text box with a fallback.",
                     "position:absolute;margin-left:0;margin-top:0;width:180pt;height:36pt;z-index:1") +
        '</mc:Fallback></mc:AlternateContent></w:r>')
    plain_vml = '<w:r>' + vml_text_box(2, "Text in a VML text box.") + '</w:r>'
    rectangle = (
        '<w:r><mc:AlternateContent><mc:Choice Requires="wps">' + wps_rectangle(3) +
        '</mc:Choice><mc:Fallback><w:pict><v:rect id="Rectangle 3" '
        'style="width:72pt;height:36pt" fillcolor="#4472c4"/></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>')
    energy = ('<m:oMathPara><m:oMath>' + m_run("E=m") + '<m:sSup><m:e>' + m_run("c") +
              '</m:e><m:sup>' + m_run("2") + '</m:sup></m:sSup></m:oMath></m:oMathPara>')
    inline = ('<m:oMath><m:f><m:num>' + m_run("a") + '</m:num><m:den>' + m_run("b") +
              '</m:den></m:f>' + m_run("+") + '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr>'
              '<m:deg/><m:e>' + m_run("z") + '</m:e></m:rad>' + m_run("+") +
              '<m:d><m:dPr><m:ctrlPr/></m:dPr><m:e>' + m_run("p+q") + '</m:e></m:d></m:oMath>')
    body = (para("Drawn text boxes and equations", "Heading1") +
            para(run("A text box with a fallback: ") + with_fallback) +
            para(run("A VML text box: ") + plain_vml) +
            para(run("A rectangle: ") + rectangle) +
            para(energy) +
            para(run("Inline: ") + inline + run(" after.")) +
            para("The end."))
    package = new_docx("Drawn content", body)
    write("features/drawn-content.docx", package.build())
    expect("features/drawn-content.docx", "docx", rtl=False, features={},
           outline=[(1, "Drawn text boxes and equations", 0)],
           text=["Drawn text boxes and equations", "A text box with a fallback:",
                 "A VML text box:", "A rectangle:", "", "Inline: after.", "The end."],
           purpose="Text boxes, a shape, and equations that docx-preview 0.4.1 draws, so the "
                   "core reports nothing as not displayed.")


def tracked_changes_fixture():
    author = 'w:author="Reviewer" w:date="2026-01-01T00:00:00Z"'
    body = (
        para("Tracked changes", "Heading1") +
        para(run("The quick ") +
             f'<w:ins w:id="1" {author}>' + run("brown ") + "</w:ins>" +
             f'<w:del w:id="2" {author}><w:r><w:delText xml:space="preserve">red </w:delText></w:r></w:del>' +
             run("fox jumps.")) +
        para(run("Second ") + f'<w:moveFromRangeStart w:id="3" {author} w:name="move1"/>' +
             f'<w:moveFrom w:id="4" {author}>' + run("alpha ") + "</w:moveFrom>" +
             '<w:moveFromRangeEnd w:id="3"/>' + run("paragraph.")) +
        para(run("Third paragraph ") + f'<w:moveToRangeStart w:id="5" {author} w:name="move1"/>' +
             f'<w:moveTo w:id="6" {author}>' + run("alpha ") + "</w:moveTo>" +
             '<w:moveToRangeEnd w:id="5"/>' + run("ends.")) +
        para('<w:r><w:rPr><w:b/><w:rPrChange w:id="7" ' + author + '><w:rPr/></w:rPrChange></w:rPr>'
             '<w:t>Formatting changed to bold.</w:t></w:r>',
             ppr=f'<w:jc w:val="center"/><w:pPrChange w:id="8" {author}><w:pPr/></w:pPrChange>') +
        '<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="6000" w:type="dxa"/></w:tblPr>'
        '<w:tblGrid><w:gridCol w:w="6000"/></w:tblGrid>'
        '<w:tr>' + cell("Kept row", 6000) + '</w:tr>'
        f'<w:tr><w:trPr><w:del w:id="9" {author}/></w:trPr>'
        f'<w:tc><w:tcPr><w:tcW w:w="6000" w:type="dxa"/></w:tcPr><w:p><w:del w:id="10" {author}>'
        '<w:r><w:delText>Deleted row</w:delText></w:r></w:del></w:p></w:tc></w:tr></w:tbl>' +
        para("End."))
    write("features/tracked-changes.docx", new_docx("Tracked changes", body).build())
    expect("features/tracked-changes.docx", "docx", rtl=False,
           features={"tracked-changes": 8},
           outline=[(1, "Tracked changes", 0)],
           text=["Tracked changes", "The quick brown fox jumps.", "Second paragraph.",
                 "Third paragraph alpha ends.", "Formatting changed to bold.", "Kept row", "End."],
           purpose="Insertions, deletions, a move, run and paragraph property changes, and a "
                   "deleted table row; the final text is known.")


def comments_fixture():
    package = new_docx("Comments", "")
    comments = XML_DECL + f'<w:comments {NS_DECLS}>' + "".join(
        f'<w:comment w:id="{index}" w:author="Reviewer" w:date="2026-01-01T00:00:00Z" w:initials="R">'
        f'{para(run(text), "CommentText")}</w:comment>'
        for index, text in ((0, "Please check this sentence."), (1, "Nice wording.")))
    comments += "</w:comments>"
    package.add("word/comments.xml", comments, CT_WML + "comments+xml")
    package.rel("word/document.xml", REL + "comments", "comments.xml")
    reference = ('<w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>'
                 '<w:commentReference w:id="{0}"/></w:r>')
    body = (para("Comments", "Heading1") +
            para('<w:commentRangeStart w:id="0"/>' + run("This sentence has a comment.") +
                 '<w:commentRangeEnd w:id="0"/>' + reference.format(0)) +
            para(run("So does ") + '<w:commentRangeStart w:id="1"/>' + run("this phrase") +
                 '<w:commentRangeEnd w:id="1"/>' + reference.format(1) + run(".")))
    package.parts["word/document.xml"] = document_xml(body)
    write("features/comments.docx", package.build())
    expect("features/comments.docx", "docx", rtl=False, features={"comments": 2},
           outline=[(1, "Comments", 0)],
           text=["Comments", "This sentence has a comment.", "So does this phrase."],
           purpose="Two comments with ranges and reference runs; the comments part is dropped.")


def long_fixture():
    words = Words(20261020)
    document = docx.Document()
    document.add_heading("A long document", 0)
    outline = []
    paragraph_index = 1
    for page in range(1, 301):
        if page % 30 == 1:
            part = (page - 1) // 30 + 1
            document.add_heading(f"Part {part}", 1)
            outline.append((1, f"Part {part}", paragraph_index))
            paragraph_index += 1
        document.add_heading(f"Section {page}", 2)
        outline.append((2, f"Section {page}", paragraph_index))
        paragraph_index += 1
        for _ in range(5):
            document.add_paragraph(words.paragraph(3))
            paragraph_index += 1
        if page < 300:
            document.add_page_break()
            paragraph_index += 1
    write("features/long-300-pages.docx", fixed_python_docx(document, "A long document"))
    # python-docx's template customXml and thumbnail are dropped, not counted.
    expect("features/long-300-pages.docx", "docx", rtl=False, outline=outline,
           purpose="300 pages built with python-docx: headings every page, a page break "
                   "after each, for the 500 ms open budget.")


# ---------------------------------------------------------------------------
# Expectations as C++


def cpp_string(value):
    out = ['"']
    for byte in value.encode("utf-8"):
        if byte == 0x22:
            out.append('\\"')
        elif byte == 0x5C:
            out.append("\\\\")
        elif byte == 0x0A:
            out.append("\\n")
        elif 0x20 <= byte < 0x7F and byte != 0x3F:
            out.append(chr(byte))
        else:
            out.append(f"\\{byte:03o}")
    out.append('"')
    return "".join(out)


def write_expectations():
    ordered = dict(sorted(EXPECTED.items()))
    (ROOT / "expected.json").write_text(json.dumps(ordered, ensure_ascii=False, indent=2) + "\n",
                                        encoding="utf-8")
    lines = ["// Generated by tests/fixtures/make-fixtures.py. Do not edit.",
             "// Fields: path, kind, ok, errorCode, features (code=count;...),",
             "// rightToLeft (-1 unchecked), outline (level|paragraph|text lines or null),",
             "// text (body paragraphs joined by newlines, or null).", ""]
    for path, item in ordered.items():
        features = ";".join(f"{code}={count}" for code, count in sorted(item["features"].items()))
        rtl = -1 if "rightToLeft" not in item else int(item["rightToLeft"])
        outline = ("nullptr" if "outline" not in item else
                   cpp_string("\n".join(f"{level}|{index}|{text}" for level, text, index in item["outline"])))
        text = "nullptr" if "text" not in item else cpp_string("\n".join(item["text"]))
        lines.append(f"{{{cpp_string(path)}, {cpp_string(item['kind'])}, "
                     f"{'true' if item['ok'] else 'false'}, {cpp_string(item['errorCode'])}, "
                     f"{cpp_string(features)}, {rtl}, {outline}, {text}}},")
    (ROOT / "expected.inc").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    security_fixtures()
    kind_fixtures()
    feature_fixtures()
    write_expectations()
    total = sum(path.stat().st_size for path in ROOT.rglob("*") if path.is_file())
    print(f"{len(EXPECTED)} fixtures, {total / 1024:.0f} KiB in tests/fixtures")


if __name__ == "__main__":
    main()
